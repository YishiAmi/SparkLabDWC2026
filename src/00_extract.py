"""Run once on the original Excel files: reads '02 Field Production.xlsx' and '01 Introduction.xlsx' into prod.pkl / unc.pkl.
Usage: python src/00_extract.py <folder containing TestCases/>   (or set DWC_RAW). Writes data/prod.pkl and data/unc.pkl"""
import openpyxl, numpy as np, pickle, sys, os
from paths import *
D=os.path.join(sys.argv[1] if len(sys.argv)>1 else RAW,"TestCases")
wb=openpyxl.load_workbook(os.path.join(D,"02 Field Production.xlsx"),read_only=True,data_only=True); out={}
for name in wb.sheetnames[2:]:
    rows=list(wb[name].iter_rows(values_only=True)); hdr=rows[1]; data=[r for r in rows[3:] if r[0] is not None]
    out[name]=([r[0] for r in data],list(hdr[1:]),np.array([[np.nan if v is None else v for v in r[1:]] for r in data],float))
pickle.dump(out,open(DATA/"prod.pkl","wb"))
wb=openpyxl.load_workbook(os.path.join(D,"01 Introduction.xlsx"),read_only=True,data_only=True)
pickle.dump(list(wb["Uncertainty Values"].iter_rows(values_only=True)),open(DATA/"unc.pkl","wb"))
